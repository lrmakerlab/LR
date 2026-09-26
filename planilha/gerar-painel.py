from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as L
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
import datetime, sys

OUT = sys.argv[1] if len(sys.argv) > 1 else "/home/user/LR/planilha/Painel-G-Code-3D.xlsx"
DEMO = len(sys.argv) > 2  # preenche dados de teste para verificação

INK = "14201A"; INK2 = "22302A"; GD = "2D7A4F"; NEON = "4ADE80"; MINT = "E6F4EA"
P2 = "F2EFE7"; RULE = "DDD8CC"; CAR = "8A5A1E"; CARS = "F6ECDC"
SOFT = "42514A"; MUTE = "6B7A72"; RED = "9C2B1E"; REDS = "F6D5D0"; WHITE = "FFFFFF"
F = "Arial"

def font(size=10, bold=False, color=INK, italic=False):
    return Font(name=F, size=size, bold=bold, color=color, italic=italic)
def fill(c): return PatternFill("solid", start_color=c, end_color=c)
thin = Side(style="thin", color=RULE); hair = Side(style="thin", color="E6E2D8")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
LINE = Border(bottom=hair)
GLINE = Border(bottom=Side(style="medium", color=NEON))
BRL = '"R$" #,##0.00;-"R$" #,##0.00;"-"'
BRL0 = '"R$" #,##0;-"R$" #,##0;"-"'
DATE = "DD/MM/YYYY"; PCT = "0%"
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
TOPL = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)

def setw(ws, widths):
    for i, w in enumerate(widths): ws.column_dimensions[L(i + 1)].width = w
def band(ws, r1, r2, c1, c2, color):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1): ws.cell(r, c).fill = fill(color)
def put(ws, ref, v, f=None, fl=None, al=None, fmt=None, border=None):
    c = ws[ref]; c.value = v
    if f: c.font = f
    if fl: c.fill = fill(fl)
    if al: c.alignment = al
    if fmt: c.number_format = fmt
    if border: c.border = border
    return c
def merge(ws, rng, v=None, f=None, fl=None, al=None, fmt=None):
    a, b = rng.split(":")
    ca, ra = coordinate_from_string(a); cb, rb = coordinate_from_string(b)
    if fl: band(ws, ra, rb, column_index_from_string(ca), column_index_from_string(cb), fl)
    ws.merge_cells(rng)
    return put(ws, a, v, f, None, al, fmt)
def header(ws, last_col, title, sub):
    ws.sheet_view.showGridLines = False
    band(ws, 1, 4, 1, last_col, INK)
    put(ws, "B2", "LR MAKER LAB · PAINEL G-CODE 3D", font(8, True, NEON))
    put(ws, "B3", title, font(20, True, WHITE))
    put(ws, "B4", sub, font(10, False, "B9C8BF"))
    ws.row_dimensions[1].height = 10; ws.row_dimensions[2].height = 16
    ws.row_dimensions[3].height = 30; ws.row_dimensions[4].height = 20
    for c in range(1, last_col + 1): ws.cell(5, c).border = GLINE
    ws.row_dimensions[5].height = 6
def section(ws, row, c1, c2, text, color=INK):
    band(ws, row, row, c1, c2, color)
    put(ws, f"{L(c1)}{row}", text, font(10, True, WHITE), al=Alignment(vertical="center", indent=1))
    ws.row_dimensions[row].height = 22
def th(ws, row, c1, labels):
    for i, lab in enumerate(labels):
        c = ws.cell(row, c1 + i, lab); c.font = font(9, True, WHITE); c.fill = fill(INK2); c.alignment = CENTER
    ws.row_dimensions[row].height = 32
def inp(c, fmt=None):
    c.fill = fill(MINT); c.font = font(10); c.border = BOX; c.alignment = LEFT
    if fmt: c.number_format = fmt
def calc(c, fmt=None, bold=False, al=None):
    c.font = font(10, bold); c.border = LINE; c.alignment = al or LEFT
    if fmt: c.number_format = fmt
def card(ws, r, c1, c2, label, formula, fmt=None, big=20, dark=False, sub=None):
    bg = INK if dark else P2
    merge(ws, f"{L(c1)}{r}:{L(c2)}{r}", label, font(8, True, NEON if dark else MUTE), bg, Alignment(horizontal="left", vertical="bottom", indent=1, wrap_text=True))
    merge(ws, f"{L(c1)}{r+1}:{L(c2)}{r+1}", formula, font(big, True, NEON if dark else INK), bg, Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True), fmt)
    merge(ws, f"{L(c1)}{r+2}:{L(c2)}{r+2}", sub, font(8, False, "B9C8BF" if dark else MUTE), bg, Alignment(horizontal="left", vertical="top", indent=1, wrap_text=True))

wb = Workbook()
SEL = "Painel!$E$7"      # G-Code escolhido
START = "Painel!$K$7"    # data do Dia 1

# =====================================================================
# BIBLIOTECA
# =====================================================================
bib = wb.active; bib.title = "Biblioteca"
G = [
 ("G-01","G-01 · Confeitarias e docerias","Nível 1","R$ 90 a 400","~R$ 65/h","terça a quinta, das 14h às 16h","as três peças em PETG","1h35 de máquina · 32 g","10 visitas → 3 propostas → 1 venda","G-04 · Restaurantes, cafés e bares"),
 ("G-02","G-02 · Nail designers e salões","Nível 1","R$ 90 a 300","~R$ 21/h","terça ou quarta, das 9h às 10h30","o mini mostruário e a plaquinha","3h00 de máquina · 60 g","10 abordagens → 3 propostas → 1 venda, mais indicações","G-05 · Semijoias e acessórios"),
 ("G-03","G-03 · Pet shops","Nível 1","R$ 160 a 800","~R$ 58/h","terça a quinta, das 14h às 16h","as tags e o mini display","2h10 de máquina · 45 g","10 visitas → 3 mini displays deixados → 1 lote em consignação","G-09 · Papelarias e festa"),
 ("G-04","G-04 · Restaurantes, cafés e bares","Nível 1","R$ 150 a 700","~R$ 52/h","terça a quinta, das 15h às 17h","o suporte, com o QR testado em 3 celulares","1h50 de máquina · 59 g","10 visitas → 3 casas com suporte em teste → 1 venda","G-01 · Confeitarias e docerias"),
 ("G-05","G-05 · Semijoias e acessórios","Nível 2","R$ 150 a 700","~R$ 25/h","terça a quinta, das 10h às 12h","o kit em PLA fosco e silk","2h40 de máquina · 60 g","10 visitas → 3 propostas → 1 venda até o próximo lançamento","G-02 · Nail designers e salões"),
 ("G-06","G-06 · Lojas de celular e assistências","Nível 2","R$ 120 a 700","~R$ 30/h","terça a quinta, das 10h às 12h","o expositor e a mini bandeja","2h20 de máquina · 65 g","10 visitas → 3 lojas medidas → 1 venda","G-10 · Oficinas e pequena indústria"),
 ("G-07","G-07 · Academias, boxes e escolas esportivas","Nível 2","R$ 250 a 1.500","~R$ 52/h","fora do pico: 10h às 12h ou 14h às 16h","a medalha com o nome do desafio deles","2h50 de máquina · 64 g","10 visitas → 3 datas de desafio anotadas → 1 pedido","G-08 · Imobiliárias e corretores"),
 ("G-08","G-08 · Imobiliárias e corretores","Nível 2","R$ 225 a 1.200","~R$ 30/h","segunda a quarta, das 9h às 11h","o kit de entrega e as tags","2h35 de máquina · 60 g","10 contatos → 3 propostas → 1 pacote, em geral de corretor autônomo","G-07 · Academias, boxes e escolas esportivas"),
 ("G-09","G-09 · Papelarias e festa","Nível 3","R$ 150 a 500","~R$ 35/h","terça ou quarta, das 10h às 11h30","o kit e o mostruário","2h05 de máquina · 33 g","10 abordagens → 2 a 3 interessadas → 1 lote-teste em até 2 semanas","G-03 · Pet shops"),
 ("G-10","G-10 · Oficinas e pequena indústria","Nível 3 · exige CAD","R$ 80 a 1.500","R$ 50 a 200/h","terça a quinta, das 14h às 16h","o gabarito, os clips e a engrenagem","3h00 de máquina · 63 g","10 visitas → 3 peças medidas → 2 pedidos","G-06 · Lojas de celular e assistências"),
]
header(bib, 12, "Biblioteca dos 10 G-Codes", "Os dados que alimentam o Painel. Vêm dos manuais e são hipóteses a validar na sua região.")
th(bib, 7, 2, ["Código", "G-Code", "Nível", "Ticket do 1º pedido", "Retorno por hora", "Quando visitar", "Kit de amostra: imprima…", "Máquina do kit", "Número honesto (expectativa, não promessa)", "Se precisar trocar"])
for i, g in enumerate(G):
    r = 8 + i
    for j, v in enumerate(g):
        c = bib.cell(r, 2 + j, v); c.font = font(9, j == 1); c.alignment = LEFT; c.border = LINE
    inp(bib.cell(r, 11)); bib.cell(r, 11).font = font(9)
    bib.row_dimensions[r].height = 32
dvb = DataValidation(type="list", formula1="=$C$8:$C$17", allow_blank=True); bib.add_data_validation(dvb); dvb.add("K8:K17")
put(bib, "B19", "A coluna \"Se precisar trocar\"", font(10, True, GD))
merge(bib, "B20:K20", "É o G-Code que o Painel sugere quando a decisão da semana 2 for TROCAR: duas levas completas, 20 visitas, sem nenhuma proposta na mão. É editável: escolha outro na lista se fizer mais sentido para você.", font(9, False, SOFT), None, LEFT)
bib.row_dimensions[20].height = 30
put(bib, "B22", "Linha do G-Code escolhido (automático, não mexa):", font(8, False, MUTE))
put(bib, "F22", f'=IFERROR(MATCH({SEL},$C$8:$C$17,0),0)', font(8, False, MUTE))
setw(bib, [2, 8, 34, 16, 16, 13, 30, 34, 20, 44, 34, 2])
bib.freeze_panes = "D8"
IDX = "Biblioteca!$F$22"
def g(col): return f'IF({IDX}=0,"",INDEX(Biblioteca!${col}$8:${col}$17,{IDX}))'

# =====================================================================
# ALVOS
# =====================================================================
alv = wb.create_sheet("Alvos")
header(alv, 18, "Alvos", "Seus comércios, da lista até a proposta na mão. Preencha no carro, logo depois de cada visita.")
R0, R1 = 13, 112
STATUS = ["A visitar", "Visitado", "Amostra deixada", "Proposta entregue", "Follow-up 48h feito", "Follow-up 7 dias feito", "Fechou", "Não agora"]
def mine(cond):  # linhas do G-Code escolhido (ou sem G-Code preenchido)
    return (f"SUMPRODUCT(((Alvos!$C${R0}:$C${R1}={SEL})+(Alvos!$C${R0}:$C${R1}=\"\"))*(Alvos!$D${R0}:$D${R1}<>\"\")*{cond})")
A_LIST = mine("1"); A_VIS = mine(f"(Alvos!$I${R0}:$I${R1}<>\"\")"); A_AMO = mine(f"(Alvos!$J${R0}:$J${R1}=\"Sim\")")
A_PROP = mine(f"(Alvos!$K${R0}:$K${R1}<>\"\")"); A_FECH = mine(f"(Alvos!$M${R0}:$M${R1}=\"Fechou\")")
A_VOLTA = mine(f"(Alvos!$P${R0}:$P${R1}<>\"\")*(Alvos!$P${R0}:$P${R1}<=TODAY())")
section(alv, 7, 2, 17, "Resumo do G-Code escolhido no Painel")
res = [("Alvos listados", A_LIST, "meta: 20"), ("Visitados", A_VIS, "meta: 10"), ("Amostras deixadas", A_AMO, "nas mãos certas"),
       ("Propostas na mão", A_PROP, "meta: 3"), ("Fecharam", A_FECH, "referência: 1"), ("Contatos para hoje", A_VOLTA, "ou atrasados")]
for i, (lab, f_, sub) in enumerate(res):
    c1 = 2 + i * 2 + (1 if i > 2 else 0)
    c1 = [2, 4, 7, 9, 12, 15][i]; c2 = c1 + 1 if i not in (1, 3, 4) else c1 + 2
    c2 = [3, 6, 8, 11, 14, 17][i]
    card(alv, 8, c1, c2, lab.upper(), "=" + f_, "0", 18, dark=(i == 3), sub=sub)
alv.row_dimensions[8].height = 18; alv.row_dimensions[9].height = 28; alv.row_dimensions[10].height = 16
cols = ["#", "G-Code", "Comércio", "Bairro", "Quem decide", "WhatsApp ou Instagram", "Nota no Google", "Data da visita", "Deixou amostra?", "Proposta entregue em", "Valor proposto", "Status", "Objeção, nas palavras exatas", "Volto em", "Próximo contato", "Oportunidade ouvida"]
th(alv, 12, 2, cols)
for r in range(R0, R1 + 1):
    c = alv.cell(r, 2, f'=IF(D{r}="","",ROW()-{R0 - 1})'); calc(c, "0", al=CENTER); c.font = font(9, False, MUTE)
    for k in range(3, 16): inp(alv.cell(r, k))
    alv.cell(r, 8).number_format = "0.0"; alv.cell(r, 9).number_format = DATE; alv.cell(r, 11).number_format = DATE
    alv.cell(r, 12).number_format = BRL0; alv.cell(r, 15).number_format = DATE
    alv.cell(r, 14).alignment = LEFT
    p = (f'=IF(OR(D{r}="",M{r}="Fechou",M{r}="Não agora"),"",IF(O{r}<>"",O{r},IF(K{r}="","",'
         f'IF(M{r}="Proposta entregue",K{r}+2,IF(M{r}="Follow-up 48h feito",K{r}+7,IF(M{r}="Follow-up 7 dias feito",K{r}+21,""))))))')
    c = alv.cell(r, 16, p); calc(c, DATE, True)
    inp(alv.cell(r, 17))
    alv.row_dimensions[r].height = 22
ex = {3: "Exemplo · apague", 4: "Confeitaria Doce Ana", 5: "Centro", 6: "Ana, a dona", 7: "@doceana", 8: 4.6,
      9: datetime.date(2026, 10, 6), 10: "Sim", 11: datetime.date(2026, 10, 6), 12: 180, 13: "Proposta entregue",
      14: "Vou ver com meu marido.", 17: "Porta-cartão com QR do Pix"}
for k, v in ex.items(): alv.cell(R0, k, v)
for k in range(3, 18): alv.cell(R0, k).font = font(10, False, MUTE, True)
dv = DataValidation(type="list", formula1=f"=Biblioteca!$C$8:$C$17", allow_blank=True); alv.add_data_validation(dv); dv.add(f"C{R0 + 1}:C{R1}")
dv2 = DataValidation(type="list", formula1='"' + ",".join(STATUS) + '"', allow_blank=True); alv.add_data_validation(dv2); dv2.add(f"M{R0}:M{R1}")
dv3 = DataValidation(type="list", formula1='"Sim,Não"', allow_blank=True); alv.add_data_validation(dv3); dv3.add(f"J{R0}:J{R1}")
alv.conditional_formatting.add(f"P{R0}:P{R1}", FormulaRule(formula=[f'AND(P{R0}<>"",P{R0}<=TODAY())'], fill=fill(NEON), font=Font(name=F, bold=True, color=INK)))
alv.conditional_formatting.add(f"B{R0}:Q{R1}", FormulaRule(formula=[f'$M{R0}="Fechou"'], fill=fill("CDEBD8")))
alv.conditional_formatting.add(f"B{R0}:Q{R1}", FormulaRule(formula=[f'$M{R0}="Não agora"'], font=Font(name=F, color="9AA39F")))
put(alv, "B11", "G-Code em branco conta para o G-Code do Painel. Status: A visitar → Visitado → Amostra deixada → Proposta entregue → Follow-up 48h → 7 dias → Fechou ou Não agora. O Próximo contato é calculado sozinho e fica verde no dia.", font(8, False, MUTE, True))
alv.row_dimensions[11].height = 16
setw(alv, [2, 5, 26, 24, 14, 18, 20, 9, 12, 11, 13, 12, 20, 32, 12, 13, 30, 2])
alv.freeze_panes = f"E{R0}"

# =====================================================================
# CALCULADORA (configuração + simulação)
# =====================================================================
cal = wb.create_sheet("Calculadora")
header(cal, 12, "Calculadora", "Faça a conta antes de responder qualquer orçamento: custo real, piso, preço sugerido e retorno por hora.")
section(cal, 7, 2, 11, "Seus números · preencha uma vez")
cfg = [("Preço do filamento", 120, BRL, "por kg. Troque pelo que você paga."),
       ("Perda e falhas", 0.15, PCT, "sobre o filamento: purga, suporte e peça perdida. Mesmo valor dos manuais."),
       ("Potência média da impressora", 150, '0" W"', "veja na etiqueta da fonte."),
       ("Tarifa de energia", 0.95, BRL, "por kWh, com impostos. Veja na conta de luz."),
       ("Piso por hora de máquina", 15, BRL, "regra do G-Code 3D: nenhum pedido abaixo disso."),
       ("Sua meta por hora", 40, BRL, "quanto você quer ganhar por hora de máquina. Define o preço sugerido.")]
for i, (lab, v, fmt, note) in enumerate(cfg):
    r = 8 + i
    merge(cal, f"B{r}:D{r}", lab, font(10), None, LEFT)
    inp(cal.cell(r, 5), fmt); cal.cell(r, 5).value = v; cal.cell(r, 5).alignment = Alignment(horizontal="right", vertical="center")
    merge(cal, f"F{r}:K{r}", note, font(9, False, MUTE, True), None, LEFT)
    cal.row_dimensions[r].height = 22
C_FIL, C_PER, C_W, C_KWH, C_PISO, C_META = [f"Calculadora!$E${8 + i}" for i in range(6)]
section(cal, 15, 2, 11, "Simule um pedido")
sim = [("O que é o pedido", "10 suportes de mesa", None), ("Filamento do pedido inteiro", 120, '0" g"'),
       ("Horas de máquina do pedido", 3.5, '0.0" h"'), ("Extras: embalagem, argola, ímã, adesivo", 8, BRL),
       ("Quantidade de peças", 10, "0"), ("Seu preço para o pedido", 200, BRL)]
for i, (lab, v, fmt) in enumerate(sim):
    r = 16 + i
    merge(cal, f"B{r}:D{r}", lab, font(10), None, LEFT)
    inp(cal.cell(r, 5), fmt); cal.cell(r, 5).value = v
    cal.row_dimensions[r].height = 22
cal["E16"].alignment = LEFT
outs = [("Filamento com perda", f"=E17/1000*{C_FIL}*(1+{C_PER})", BRL),
        ("Energia", f"=E18*{C_W}/1000*{C_KWH}", BRL),
        ("Custo real", "=H16+H17+E19", BRL),
        ("Piso (custo + horas × piso)", f"=H18+E18*{C_PISO}", BRL),
        ("Preço sugerido (custo + horas × meta)", f"=H18+E18*{C_META}", BRL),
        ("Margem", "=E21-H18", BRL),
        ("Retorno por hora de máquina", '=IF(E18=0,"",H21/E18)', BRL),
        ("Preço por peça", '=IF(E20=0,"",E21/E20)', BRL)]
for i, (lab, f_, fmt) in enumerate(outs):
    r = 16 + i
    merge(cal, f"F{r}:G{r}", lab, font(10, i in (2, 6), SOFT), None, LEFT)
    c = cal.cell(r, 8, f_); calc(c, fmt, i in (2, 6), Alignment(horizontal="right", vertical="center"))
    cal.row_dimensions[r].height = 22
merge(cal, "B25:K26", f'=IF(E21="","",IF(H22<{C_PISO},"ABAIXO DO PISO. Não feche por esse preço: suba o preço ou mude a peça.",IF(H22<{C_META},"Acima do piso e abaixo da sua meta. Dá para fechar, mas veja se cabe subir o preço.","Na meta. Pode fechar.")))',
      font(12, True, INK), P2, LEFT)
cal.conditional_formatting.add("B25", FormulaRule(formula=[f'$H$22<{C_PISO}'], fill=fill(REDS), font=Font(name=F, bold=True, size=12, color=RED)))
cal.conditional_formatting.add("B25", FormulaRule(formula=[f'$H$22>={C_META}'], fill=fill(NEON), font=Font(name=F, bold=True, size=12, color=INK)))
merge(cal, "B28:K29", "Horas de máquina: use o tempo do fatiador para o pedido inteiro. O tempo de arte, mensagem e ida à loja não entra aqui; é por causa dele que o piso existe. Todo número é hipótese a validar com a sua impressora e a sua região.", font(9, False, MUTE, True), None, LEFT)
setw(cal, [2, 14, 14, 14, 16, 18, 22, 16, 10, 10, 10, 2])

# =====================================================================
# RESULTADOS (vendas reais)
# =====================================================================
rs = wb.create_sheet("Resultados")
header(rs, 17, "Resultados", "Cada venda fechada, com o custo real e o retorno por hora de verdade. Usa os seus números da Calculadora.")
S0, S1 = 9, 108
th(rs, 8, 2, ["Data", "G-Code", "Cliente", "O que vendeu", "Qtd", "Valor recebido", "Filamento (g)", "Horas de máquina", "Extras", "Custo real", "Lucro", "Margem", "Retorno por hora", "Situação", "Pediu de novo?"])
put(rs, "B7", "G-Code em branco conta para o G-Code do Painel. O retorno por hora fica vermelho abaixo do piso.", font(8, False, MUTE, True))
for r in range(S0, S1 + 1):
    for k in range(2, 11): inp(rs.cell(r, k))
    rs.cell(r, 2).number_format = DATE; rs.cell(r, 6).number_format = "0"; rs.cell(r, 7).number_format = BRL
    rs.cell(r, 8).number_format = "0"; rs.cell(r, 9).number_format = "0.0"; rs.cell(r, 10).number_format = BRL
    calc(rs.cell(r, 11, f'=IF(G{r}="","",H{r}/1000*{C_FIL}*(1+{C_PER})+I{r}*{C_W}/1000*{C_KWH}+J{r})'), BRL)
    calc(rs.cell(r, 12, f'=IF(G{r}="","",G{r}-K{r})'), BRL, True)
    calc(rs.cell(r, 13, f'=IF(OR(G{r}="",G{r}=0),"",L{r}/G{r})'), PCT)
    calc(rs.cell(r, 14, f'=IF(OR(G{r}="",I{r}="",I{r}=0),"",L{r}/I{r})'), BRL, True)
    calc(rs.cell(r, 15, f'=IF(N{r}="","",IF(N{r}<{C_PISO},"Abaixo do piso",IF(N{r}<{C_META},"Acima do piso","Na meta")))'))
    inp(rs.cell(r, 16))
    rs.row_dimensions[r].height = 22
exr = {2: datetime.date(2026, 10, 14), 3: "Exemplo · apague", 4: "Confeitaria Doce Ana", 5: "Kit de 3 cortadores e carimbo", 6: 3,
       7: 180, 8: 90, 9: 2.5, 10: 6, 16: "Sim"}
for k, v in exr.items(): rs.cell(S0, k, v)
for k in range(2, 17): rs.cell(S0, k).font = font(10, k in (12, 14), MUTE, True)
dvr = DataValidation(type="list", formula1="=Biblioteca!$C$8:$C$17", allow_blank=True); rs.add_data_validation(dvr); dvr.add(f"C{S0 + 1}:C{S1}")
dvr2 = DataValidation(type="list", formula1='"Sim,Não,Ainda não"', allow_blank=True); rs.add_data_validation(dvr2); dvr2.add(f"P{S0}:P{S1}")
rs.conditional_formatting.add(f"N{S0}:O{S1}", FormulaRule(formula=[f'AND($N{S0}<>"",$N{S0}<{C_PISO})'], fill=fill(REDS), font=Font(name=F, bold=True, color=RED)))
rs.conditional_formatting.add(f"N{S0}:O{S1}", FormulaRule(formula=[f'AND($N{S0}<>"",$N{S0}>={C_META})'], fill=fill("CDEBD8"), font=Font(name=F, bold=True, color=GD)))
setw(rs, [2, 12, 22, 22, 26, 7, 14, 12, 11, 11, 12, 12, 9, 13, 14, 13, 2])
rs.freeze_panes = f"E{S0}"
def rmine(cond):
    return f"SUMPRODUCT(((Resultados!$C${S0}:$C${S1}={SEL})+(Resultados!$C${S0}:$C${S1}=\"\"))*(Resultados!$G${S0}:$G${S1}<>\"\")*{cond})"
R_FAT = rmine(f"N(+Resultados!$G${S0}:$G${S1})") if False else rmine(f"IFERROR(Resultados!$G${S0}:$G${S1}*1,0)")
R_LUC = rmine(f"IFERROR(Resultados!$L${S0}:$L${S1}*1,0)")
R_HOR = rmine(f"IFERROR(Resultados!$I${S0}:$I${S1}*1,0)")
R_N = rmine("1")
R_BAIXO = rmine(f"IFERROR((Resultados!$N${S0}:$N${S1}*1)<{C_PISO},0)")

# =====================================================================
# EXPANSÃO
# =====================================================================
ex_ = wb.create_sheet("Expansão")
header(ex_, 13, "Expansão", "A primeira venda abre a porta. Na entrega, faça as três perguntas da seção 11 do manual e anote tudo aqui.")
E0, E1 = 12, 111
merge(ex_, "B7:L8", "Meta: três oportunidades por cliente atendido. Quem já comprou de você e recebeu no prazo compra de novo sem visita fria. Cliente que ainda não tem três linhas aqui ainda tem venda escondida.", font(10, True, INK), CARS, LEFT)
ex_.row_dimensions[7].height = 20; ex_.row_dimensions[8].height = 20
th(ex_, 11, 2, ["Data", "Cliente", "G-Code", "Pergunta", "O que ele respondeu", "Oportunidade", "Valor estimado", "Próxima ação", "Prazo", "Status", "Oportunidades deste cliente"])
PERG = ["Pergunta 1 da seção 11", "Pergunta 2 da seção 11", "Pergunta 3 da seção 11", "Vi na operação", "Indicação"]
EST = ["Anotada", "Proposta feita", "Fechou", "Não agora"]
for r in range(E0, E1 + 1):
    for k in range(2, 12): inp(ex_.cell(r, k))
    ex_.cell(r, 2).number_format = DATE; ex_.cell(r, 8).number_format = BRL0; ex_.cell(r, 10).number_format = DATE
    for k in (6, 7, 9): ex_.cell(r, k).alignment = LEFT
    calc(ex_.cell(r, 12, f'=IF(C{r}="","",COUNTIF($C${E0}:$C${E1},C{r}))'), "0", True, CENTER)
    ex_.row_dimensions[r].height = 26
exe = {2: datetime.date(2026, 10, 20), 3: "Exemplo · apague", 5: "Indicação", 6: "\"A gente compra sobremesa da confeitaria da esquina.\"",
       7: "Indicação: confeitaria parceira", 8: 150, 9: "Pedir o contato e levar o kit", 10: datetime.date(2026, 10, 24), 11: "Anotada"}
for k, v in exe.items(): ex_.cell(E0, k, v)
for k in range(2, 12): ex_.cell(E0, k).font = font(10, False, MUTE, True)
dve = DataValidation(type="list", formula1="=Biblioteca!$C$8:$C$17", allow_blank=True); ex_.add_data_validation(dve); dve.add(f"D{E0}:D{E1}")
dve2 = DataValidation(type="list", formula1='"' + ",".join(PERG) + '"', allow_blank=True); ex_.add_data_validation(dve2); dve2.add(f"E{E0}:E{E1}")
dve3 = DataValidation(type="list", formula1='"' + ",".join(EST) + '"', allow_blank=True); ex_.add_data_validation(dve3); dve3.add(f"K{E0}:K{E1}")
ex_.conditional_formatting.add(f"L{E0}:L{E1}", FormulaRule(formula=[f'AND(L{E0}<>"",L{E0}>=3)'], fill=fill(NEON), font=Font(name=F, bold=True, color=INK)))
ex_.conditional_formatting.add(f"B{E0}:L{E1}", FormulaRule(formula=[f'$K{E0}="Fechou"'], fill=fill("CDEBD8")))
setw(ex_, [2, 12, 22, 22, 22, 32, 28, 13, 26, 12, 14, 14, 2])
ex_.freeze_panes = f"D{E0}"
X_N = f"SUMPRODUCT(((Expansão!$D${E0}:$D${E1}={SEL})+(Expansão!$D${E0}:$D${E1}=\"\"))*(Expansão!$C${E0}:$C${E1}<>\"\")*(Expansão!$C${E0}:$C${E1}<>\"Exemplo · apague\"))"
X_F = f"SUMPRODUCT(((Expansão!$D${E0}:$D${E1}={SEL})+(Expansão!$D${E0}:$D${E1}=\"\"))*(Expansão!$C${E0}:$C${E1}<>\"\")*(Expansão!$K${E0}:$K${E1}=\"Fechou\"))"

# =====================================================================
# MISSÃO 72H
# =====================================================================
mi = wb.create_sheet("Missão 72h")
header(mi, 11, "Missão 72 horas", "Em 72 horas, você sai da impressora parada para três propostas entregues na mão de um comerciante.")
M = [  # (bloco, dia, ação, como saber, números)
 ("DIA 1 · PREPARAR", 1, '="Imprimir dois jogos do kit de amostra: "&' + g("H") + '&IF(' + IDX + '=0,""," ("&' + g("I") + '&")")', "Kit conferido, sem defeito, seção 06 do manual", '""'),
 ("", 1, '="Montar a embalagem e imprimir a proposta com os três níveis de preço"', "Três propostas impressas, com preço em cada nível", '""'),
 ("", 1, '="Listar 20 alvos na aba Alvos, com os sinais da seção 02 do manual"', "20 linhas na aba Alvos", f'=({A_LIST})&" de 20 alvos"'),
 ("", 1, '="Decorar o script de 40 segundos e as seis respostas de objeção"', "Você fala em voz alta sem ler", '""'),
 ("DIA 2 · RUA", 2, '="Visitar 10 alvos"&IF(' + IDX + '=0,""," · "&' + g("G") + ')', "10 datas de visita na aba Alvos", f'=({A_VIS})&" de 10 visitas"'),
 ("", 2, '="Entregar 3 propostas na mão de quem decide"', "3 datas em Proposta entregue", f'=({A_PROP})&" de 3 propostas"'),
 ("", 2, '="Anotar toda objeção com as palavras exatas, no carro"', "Coluna Objeção preenchida", '""'),
 ("DIA 3 · FECHAR", 3, '="Mandar o follow-up para as três propostas (seção 10)"', "Status Follow-up 48h feito", '""'),
 ("", 3, '="Registrar tudo na aba Alvos, com a data de volta de cada um"', "Nenhuma visita só no papel", '""'),
 ("", 3, '="Anotar as oportunidades ouvidas no balcão"', "Coluna Oportunidade ouvida", '""'),
 ("", 3, '="Marcar na agenda a data da segunda leva de 10 visitas"', "Data escrita na agenda", '""'),
 ("SEMANA 2 · CONTINUAR", 4, '="Voltar nos alvos com amostra ou proposta, no dia combinado"', "Próximo contato em dia na aba Alvos", f'=({A_VOLTA})&" contatos para hoje ou atrasados"'),
 ("", 4, '="Fazer a segunda leva de 10 visitas, já com as objeções reais anotadas"', "20 visitas no total", f'=({A_VIS})&" de 20 visitas"'),
 ("", 4, '="Produzir o primeiro pedido e conferir com a lista da seção 05"', "Venda registrada na aba Resultados", f'=({R_N})&" vendas registradas"'),
 ("", 4, '="Na entrega, fazer as três perguntas da seção 11"', "Linhas na aba Expansão", f'=({X_N})&" oportunidades anotadas"'),
]
th(mi, 7, 2, ["Etapa", "Quando", "O que fazer", "Como saber que está feito", "Os números dizem", "Feito?", "Situação"])
mi.merge_cells("D7:E7") if False else None
r = 8; M0 = 8
for bloco, dia, acao, como, num in M:
    if bloco:
        section(mi, r, 2, 10, bloco, INK if dia < 4 else GD); r += 1
    put(mi, f"B{r}", "", font(9))
    c = mi.cell(r, 3, f'=IF({START}="","Dia {dia}",{START}+{dia - 1})' if dia < 4 else f'=IF({START}="","Dias 4 a 10",{START}+3)')
    calc(c, DATE, al=CENTER); c.font = font(9, False, SOFT)
    c = mi.cell(r, 4, acao); calc(c); c.font = font(10, True)
    merge(mi, f"E{r}:F{r}", como, font(9, False, SOFT), None, LEFT)
    c = mi.cell(r, 7, None if num == '""' else num); calc(c); c.font = font(9, True, GD)
    inp(mi.cell(r, 8)); mi.cell(r, 8).value = "Não"; mi.cell(r, 8).alignment = CENTER
    st = (f'=IF(H{r}="Sim","Feito",IF({START}="","A fazer",IF({START}+{min(dia, 4) - 1}<TODAY(),"Atrasado",'
          f'IF({START}+{min(dia, 4) - 1}=TODAY(),"Hoje","A fazer"))))')
    c = mi.cell(r, 9, st); calc(c, al=CENTER); c.font = font(9, True)
    mi.cell(r, 2).value = dia
    mi.cell(r, 2).font = font(9, True, "B9C8BF"); mi.cell(r, 2).alignment = CENTER
    mi.row_dimensions[r].height = 34
    r += 1
M_LAST = r - 1
DONE72 = "SUMPRODUCT(('Missão 72h'!$B$8:$B$" + str(M_LAST) + "<4)*('Missão 72h'!$B$8:$B$" + str(M_LAST) + "<>\"\")*('Missão 72h'!$H$8:$H$" + str(M_LAST) + "=\"Sim\"))"
TOT72 = "SUMPRODUCT(('Missão 72h'!$B$8:$B$" + str(M_LAST) + "<4)*('Missão 72h'!$B$8:$B$" + str(M_LAST) + "<>\"\")*1)"
dvm = DataValidation(type="list", formula1='"Sim,Não"', allow_blank=False); mi.add_data_validation(dvm); dvm.add(f"H8:H{M_LAST}")
mi.conditional_formatting.add(f"I8:I{M_LAST}", FormulaRule(formula=['I8="Feito"'], fill=fill(NEON), font=Font(name=F, bold=True, color=INK)))
mi.conditional_formatting.add(f"I8:I{M_LAST}", FormulaRule(formula=['I8="Hoje"'], fill=fill(INK), font=Font(name=F, bold=True, color=NEON)))
mi.conditional_formatting.add(f"I8:I{M_LAST}", FormulaRule(formula=['I8="Atrasado"'], fill=fill(REDS), font=Font(name=F, bold=True, color=RED)))
mi.conditional_formatting.add(f"H8:H{M_LAST}", FormulaRule(formula=['H8="Sim"'], fill=fill("CDEBD8"), font=Font(name=F, bold=True, color=GD)))
# esconder a coluna B (número do dia, usada nas contas) visualmente
mi.column_dimensions["B"].width = 3
r = M_LAST + 2
DEC_R = r
section(mi, r, 2, 10, "DECISÃO DA SEMANA 2 · o Painel lê os seus números")
DEC = (f'=IF({IDX}=0,"ESCOLHA SEU G-CODE",IF(({A_FECH})+({R_N})>=1,"ESCALAR",IF(AND(({A_VIS})>=20,({A_PROP})=0),"TROCAR",'
       f'IF(({A_PROP})>=3,"AJUSTAR O FOLLOW-UP",IF(({A_VIS})>=10,"AJUSTAR A ABORDAGEM","SIGA A MISSÃO")))))')
DEC_TXT = (f'=IF({IDX}=0,"Escolha seu G-Code no Painel.",IF(({A_FECH})+({R_N})>=1,"Houve venda. Produza, entregue no prazo, faça as três perguntas da seção 11 e faça a segunda leva de visitas no mesmo bairro.",'
           f'IF(AND(({A_VIS})>=20,({A_PROP})=0),"Duas levas completas sem nenhuma proposta. Troque para "&{g("K")}&". Leve as objeções que você anotou: elas valem para o próximo.",'
           f'IF(({A_PROP})>=3,"As propostas estão na mão. Agora é follow-up: 48h, 7 dias e a mensagem final, como na seção 10. Não troque de G-Code antes disso.",'
           f'IF(({A_VIS})>=10,"Você visitou, mas ainda não chegou a 3 propostas. Releia as objeções anotadas, confira o horário de visita e faça a segunda leva de 10.",'
           f'"Siga a próxima ação do Painel. A decisão fica valendo depois das primeiras 10 visitas.")))))')
merge(mi, f"B{r+1}:C{r+2}", DEC, font(14, True, NEON), INK, CENTER)
merge(mi, f"D{r+1}:J{r+2}", DEC_TXT, font(10, False, INK), P2, LEFT)
mi.row_dimensions[r + 1].height = 26; mi.row_dimensions[r + 2].height = 26
for i, (t, d) in enumerate([("ESCALAR", "Houve venda: repita no mesmo bairro."), ("AJUSTAR", "Propostas na mão ou visitas sem proposta: ajuste antes de trocar."), ("TROCAR", "Só depois de 20 visitas sem nenhuma proposta.")]):
    c1 = [2, 5, 8][i]; c2 = [4, 7, 10][i]
    merge(mi, f"{L(c1)}{r+4}:{L(c2)}{r+4}", t, font(9, True, GD), None, LEFT)
    merge(mi, f"{L(c1)}{r+5}:{L(c2)}{r+5}", d, font(9, False, SOFT), None, LEFT)
mi.row_dimensions[r + 5].height = 28
setw(mi, [2, 3, 12, 58, 22, 22, 26, 9, 12, 2, 2])
mi.freeze_panes = "A8"

# =====================================================================
# PAINEL
# =====================================================================
pa = wb.create_sheet("Painel", 0)
header(pa, 14, "Painel", "Onde você está, o que fazer hoje e o que os seus números dizem.")
put(pa, "B7", "Seu G-Code", font(9, True, MUTE), al=Alignment(vertical="center"))
merge(pa, "E7:H7", None)
for c in range(5, 9): inp(pa.cell(7, c))
pa["E7"].font = font(12, True, INK)
put(pa, "J7", "Dia 1 da missão", font(9, True, MUTE), al=Alignment(horizontal="right", vertical="center"))
inp(pa["K7"], DATE); pa["K7"].font = font(12, True, INK); pa["K7"].alignment = CENTER
merge(pa, "L7:M7", '=IF(K7="","← data do Dia 1","")', font(8, False, MUTE, True), None, LEFT)
merge(pa, "B7:D7", "Seu G-Code (o que o Seletor indicou)", font(9, True, MUTE), None, Alignment(vertical="center"))
pa.row_dimensions[7].height = 30
dvp = DataValidation(type="list", formula1="=Biblioteca!$C$8:$C$17", allow_blank=True); pa.add_data_validation(dvp); dvp.add("E7")
put(pa, "B8", '=IF(E7="","Escolha seu G-Code na lista acima. O Painel inteiro se ajusta a ele.",' + g("D") + '&" · ticket do 1º pedido "&' + g("E") + '&" · "&' + g("F") + '&" · visite "&' + g("G") + ')', font(9, False, SOFT, True))
pa.merge_cells("B8:M8"); pa.row_dimensions[8].height = 18

DIA = f'IF({START}="","",TODAY()-{START}+1)'
PROG = f"({DONE72})/({TOT72})"
card(pa, 10, 2, 4, "DIA DA MISSÃO", f'=IF(K7="","—",IF(TODAY()<K7,"Começa em "&(K7-TODAY())&" dia(s)",IF(TODAY()-K7+1<=3,"Dia "&(TODAY()-K7+1)&" de 3","Semana 2")))', None, 18, True, "72 horas: preparar, rua, fechar")
card(pa, 10, 5, 7, "MISSÃO 72H CONCLUÍDA", f"={PROG}", PCT, 22, False, f'=REPT("■",ROUND(({PROG})*20,0))&REPT("□",20-ROUND(({PROG})*20,0))')
card(pa, 10, 8, 10, "PROPOSTAS NA MÃO", f'=({A_PROP})&" de 3"', None, 22, False, "a promessa do G-Code 3D")
card(pa, 10, 11, 13, "ONDE VOCÊ ESTÁ", f'=IF(E7="","—",IF(AND(({A_FECH})+({R_N})>=1,({X_N})>=1),"Expandindo",IF(({A_FECH})+({R_N})>=1,"Vendendo",IF(({A_PROP})>=3,"Missão cumprida",IF(({A_PROP})>=1,"Com propostas",IF(({A_VIS})>=1,"Na rua",IF(({A_LIST})>=1,"Mapeando","Preparando")))))))', None, 16, False, "Preparando → Na rua → Propostas → Vendendo → Expandindo")
for rr in (10, 11, 12): pa.row_dimensions[rr].height = [18, 34, 18][rr - 10]

# próxima ação
MR = f"'Missão 72h'!$H$8:$H${M_LAST}"
NEXT = (f'=IF(E7="","Escolha seu G-Code acima e marque a data do Dia 1.",IF(K7="","Marque a data do Dia 1 ao lado do G-Code.",'
        f'IFERROR("Dia "&INDEX(\'Missão 72h\'!$B$8:$B${M_LAST},MATCH("Não",{MR},0))&": "&INDEX(\'Missão 72h\'!$D$8:$D${M_LAST},MATCH("Não",{MR},0)),'
        f'"Tudo marcado. Leia a decisão da semana 2 abaixo.")))')
section(pa, 14, 2, 13, "PRÓXIMA AÇÃO")
merge(pa, "B15:M16", NEXT, font(14, True, INK), P2, LEFT)
pa.row_dimensions[15].height = 24; pa.row_dimensions[16].height = 24
merge(pa, "B17:M17", "A primeira ação da Missão 72h que ainda está como \"Não\". Marque \"Sim\" na aba Missão 72h quando fizer.", font(8, False, MUTE, True), None, LEFT)

# funil
section(pa, 19, 2, 7, "FUNIL DO SEU G-CODE")
th(pa, 20, 2, ["Etapa", "Feito", "Meta", "Progresso", "", ""])
pa.merge_cells("E20:G20")
fun = [("Alvos listados", A_LIST, 20), ("Visitados", A_VIS, 10), ("Amostras deixadas", A_AMO, 3), ("Propostas na mão", A_PROP, 3), ("Fecharam", A_FECH, 1)]
for i, (lab, f_, meta) in enumerate(fun):
    r = 21 + i
    c = pa.cell(r, 2, lab); calc(c); c.font = font(10, i == 3)
    c = pa.cell(r, 3, "=" + f_); calc(c, "0", True, CENTER)
    c = pa.cell(r, 4, meta); calc(c, "0", False, CENTER); c.font = font(10, False, MUTE)
    merge(pa, f"E{r}:G{r}", f'=REPT("■",MIN(10,ROUND(C{r}/D{r}*10,0)))&REPT("□",10-MIN(10,ROUND(C{r}/D{r}*10,0)))', font(11, False, GD), None, LEFT)
    pa.row_dimensions[r].height = 22
merge(pa, "B26:G26", "Referência de Fecharam: 1 venda a cada 10 visitas é o número honesto dos manuais. É expectativa, não promessa: a venda depende de quem decide do outro lado do balcão.", font(8, False, MUTE, True), None, LEFT)
pa.row_dimensions[26].height = 26

# resumo do G-Code
section(pa, 19, 9, 13, "SEU G-CODE EM UMA OLHADA")
info = [("Nível", g("D")), ("Ticket do 1º pedido", g("E")), ("Retorno por hora", g("F")), ("Quando visitar", g("G")),
        ("Kit de amostra", g("H")), ("Número honesto", g("J")), ("Manual", f'IF({IDX}=0,"",INDEX(Biblioteca!$B$8:$B$17,{IDX})&" · leia a Cola de balcão antes de entrar")')]
for i, (lab, f_) in enumerate(info):
    r = 20 + i
    c = pa.cell(r, 9, lab); c.font = font(9, True, MUTE); c.alignment = LEFT; c.border = LINE
    merge(pa, f"J{r}:M{r}", "=" + f_, font(10, i in (1, 2)), None, LEFT)
    for k in range(10, 14): pa.cell(r, k).border = LINE
    pa.row_dimensions[r].height = 22 if i < 4 else 30

# resultados
section(pa, 28, 2, 13, "RESULTADOS · da aba Resultados")
card(pa, 29, 2, 4, "VENDAS", "=" + R_N, "0", 20)
card(pa, 29, 5, 7, "FATURAMENTO", "=" + R_FAT, BRL0, 20)
card(pa, 29, 8, 10, "LUCRO", "=" + R_LUC, BRL0, 20)
card(pa, 29, 11, 13, "LUCRO POR HORA DE MÁQUINA", f'=IF(({R_HOR})=0,"—",({R_LUC})/({R_HOR}))', BRL, 20, True, f'=IF(({R_BAIXO})>0,({R_BAIXO})&" venda(s) abaixo do piso","nenhuma venda abaixo do piso")')
pa.row_dimensions[30].height = 34
section(pa, 33, 2, 13, "EXPANSÃO · da aba Expansão")
card(pa, 34, 2, 4, "OPORTUNIDADES ANOTADAS", "=" + X_N, "0", 20)
card(pa, 34, 5, 7, "OPORTUNIDADES FECHADAS", "=" + X_F, "0", 20)
merge(pa, "H34:M36", "Meta: três oportunidades por cliente atendido. Faça as três perguntas da seção 11 na entrega de todo pedido.", font(10, False, SOFT), P2, LEFT)
pa.row_dimensions[35].height = 34

# decisão
section(pa, 38, 2, 13, "DECISÃO DA SEMANA 2")
merge(pa, "B39:D40", f"='Missão 72h'!B{DEC_R + 1}", font(14, True, NEON), INK, CENTER)
merge(pa, "E39:M40", f"='Missão 72h'!D{DEC_R + 1}", font(10, False, INK), P2, LEFT)
pa.row_dimensions[39].height = 26; pa.row_dimensions[40].height = 26
merge(pa, "B42:M42", "Todo número deste painel é hipótese a validar com a sua impressora, o seu filamento e a sua região. LR Maker Lab · G-Code 3D · uso pessoal do comprador.", font(8, False, MUTE, True), None, LEFT)
pa.conditional_formatting.add("B11", FormulaRule(formula=['$K$7=""'], font=Font(name=F, bold=True, size=18, color="9AA39F")))
setw(pa, [2, 16, 10, 9, 10, 10, 12, 11, 20, 14, 14, 14, 16, 2])

# =====================================================================
# COMECE AQUI
# =====================================================================
ca = wb.create_sheet("Comece Aqui", 0)
header(ca, 12, "Comece aqui", "O manual diz o que fazer. Este painel mostra se você está fazendo.")
steps = [("1", "Descubra o seu G-Code", "Responda as três perguntas do Seletor, na sua área de membros. Ele indica um G-Code só. Os outros nove ficam na biblioteca, para depois da primeira venda."),
         ("2", "Escolha no Painel", "Na aba Painel, escolha o G-Code na lista e marque a data do Dia 1. Todas as abas se ajustam a ele."),
         ("3", "Faça a Missão 72h", "Dia 1 preparar, Dia 2 rua, Dia 3 fechar. Marque \"Sim\" em cada ação feita. O Painel mostra sempre a próxima."),
         ("4", "Registre tudo", "Visitas e propostas na aba Alvos, no carro, logo depois de cada visita. Vendas na aba Resultados, com as horas de máquina."),
         ("5", "Expanda no cliente", "Na entrega do primeiro pedido, faça as três perguntas da seção 11 e anote na aba Expansão. Meta: três oportunidades por cliente.")]
r = 7
for n, t, d in steps:
    merge(ca, f"B{r}:B{r+1}", n, font(24, True, INK), NEON, CENTER)
    merge(ca, f"C{r}:K{r}", t, font(12, True, INK), P2, Alignment(vertical="bottom", indent=1))
    merge(ca, f"C{r+1}:K{r+1}", d, font(10, False, SOFT), None, Alignment(vertical="top", wrap_text=True, indent=1))
    ca.row_dimensions[r].height = 22; ca.row_dimensions[r + 1].height = 34
    r += 3
section(ca, r, 2, 11, "A REGRA DESTE PAINEL")
merge(ca, f"B{r+1}:K{r+2}", "Missão cumprida é três propostas entregues na mão de um comerciante. A venda depende de quem decide do outro lado do balcão, e por isso não é a meta da missão: é o que vem depois de visitas bem feitas e follow-up no prazo.", font(10, True, INK), MINT, LEFT)
ca.row_dimensions[r + 1].height = 26; ca.row_dimensions[r + 2].height = 26
r += 4
section(ca, r, 2, 11, "AS CORES")
leg = [(MINT, INK, "Você preenche"), (None, INK, "Calculado sozinho, não mexa"), (NEON, INK, "Pronto, ou para hoje"), (REDS, RED, "Atenção: atrasado ou abaixo do piso")]
for i, (bg, fg, t) in enumerate(leg):
    rr = r + 1 + i
    c = ca.cell(rr, 2, "Aa"); c.font = font(10, True, fg); c.alignment = CENTER; c.border = BOX
    if bg: c.fill = fill(bg)
    merge(ca, f"C{rr}:K{rr}", t, font(10, False, SOFT), None, Alignment(vertical="center", indent=1))
    ca.row_dimensions[rr].height = 20
r += 6
section(ca, r, 2, 11, "AS ABAS")
abas = [("Painel", "onde você está, a próxima ação e a decisão da semana 2"), ("Missão 72h", "as ações dos três dias e da semana 2, com data"),
        ("Alvos", "seus comércios, da lista até a proposta na mão"), ("Calculadora", "a conta do preço antes de responder um orçamento"),
        ("Resultados", "cada venda, com lucro e retorno por hora de verdade"), ("Expansão", "as oportunidades dentro de quem já comprou"),
        ("Biblioteca", "os dados dos 10 G-Codes que alimentam o Painel")]
for i, (a, d) in enumerate(abas):
    rr = r + 1 + i
    c = ca.cell(rr, 2, a); c.font = font(10, True, GD); c.border = LINE
    ca.merge_cells(f"B{rr}:C{rr}")
    merge(ca, f"D{rr}:K{rr}", d, font(10, False, SOFT), None, Alignment(vertical="center", indent=1))
    ca.row_dimensions[rr].height = 20
setw(ca, [2, 10, 12, 12, 12, 12, 12, 12, 12, 12, 12, 2])

# ordem e aparência
order = ["Comece Aqui", "Painel", "Missão 72h", "Alvos", "Calculadora", "Resultados", "Expansão", "Biblioteca"]
wb._sheets = [wb[n] for n in order]
tabs = {"Comece Aqui": NEON, "Painel": INK, "Missão 72h": INK, "Alvos": GD, "Calculadora": GD, "Resultados": GD, "Expansão": GD, "Biblioteca": "9AA39F"}
for ws in wb.worksheets:
    ws.sheet_properties.tabColor = tabs[ws.title]
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.active = 0

if DEMO:  # dados de teste só na cópia de verificação
    pa["E7"] = "G-04 · Restaurantes, cafés e bares"; pa["K7"] = datetime.date.today() - datetime.timedelta(days=1)
    for i in range(12):
        r = R0 + 1 + i
        alv.cell(r, 4, f"Casa {i+1}"); alv.cell(r, 5, "Centro")
        if i < 8: alv.cell(r, 9, datetime.date.today() - datetime.timedelta(days=1))
        if i < 3:
            alv.cell(r, 10, "Sim"); alv.cell(r, 11, datetime.date.today() - datetime.timedelta(days=2)); alv.cell(r, 13, "Proposta entregue")
    for k in range(3): mi.cell(M0 + 1 + k, 8, "Sim")
    rs.cell(S0 + 1, 2, datetime.date.today()); rs.cell(S0 + 1, 4, "Casa 1"); rs.cell(S0 + 1, 6, 10); rs.cell(S0 + 1, 7, 200); rs.cell(S0 + 1, 8, 120); rs.cell(S0 + 1, 9, 3.5); rs.cell(S0 + 1, 10, 8)
    rs.cell(S0 + 2, 2, datetime.date.today()); rs.cell(S0 + 2, 4, "Casa 2"); rs.cell(S0 + 2, 7, 60); rs.cell(S0 + 2, 8, 200); rs.cell(S0 + 2, 9, 6); rs.cell(S0 + 2, 10, 5)

wb.save(OUT)
print("ok", OUT, "missão até linha", M_LAST, "decisão", DEC_R)
